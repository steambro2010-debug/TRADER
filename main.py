import os
import re
import sys
import subprocess
from pathlib import Path
from datetime import datetime

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from PySide6.QtCore import Qt, QDate
from PySide6.QtGui import QFont
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QDateEdit,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QStackedWidget,
    QVBoxLayout,
    QWidget,
)

CLINIC_NAME = "Malathi Dental Clinic"
DOCTOR_NAME = "Dr. Malathi Thandapani, BDS"
TREATMENTS = [
    "Consultation",
    "Scaling",
    "Tooth Extraction",
    "Root Canal Treatment",
    "Dental Filling",
    "Crown Fixing",
    "Teeth Whitening",
    "Braces Consultation",
    "Implant Consultation",
    "Pediatric Dental Checkup",
]
PAYMENT_METHODS = ["Cash", "GPay", "PhonePe", "Card"]


class FileManager:
    def __init__(self, base_dir: Path) -> None:
        self.base_dir = base_dir
        self.data_dir = self.base_dir / "data"
        self.patients_dir = self.data_dir / "Patients"
        self.records_file = self.data_dir / "MASTER_RECORDS.docx"
        self._ensure_structure()

    def _ensure_structure(self) -> None:
        self.patients_dir.mkdir(parents=True, exist_ok=True)
        if not self.records_file.exists():
            doc = Document()
            doc.add_heading(f"{CLINIC_NAME} - Master Records", level=1)
            doc.add_paragraph(f"Doctor: {DOCTOR_NAME}")
            doc.add_paragraph("\n")
            doc.save(self.records_file)

    @staticmethod
    def clean_patient_name(name: str) -> str:
        cleaned = re.sub(r"[^A-Za-z0-9 _-]", "", name).strip()
        cleaned = re.sub(r"\s+", "_", cleaned)
        return cleaned or "Unknown_Patient"

    def patient_dir(self, patient_name: str) -> Path:
        folder_name = self.clean_patient_name(patient_name)
        path = self.patients_dir / folder_name
        path.mkdir(parents=True, exist_ok=True)
        return path

    def generate_bill_pdf(
        self,
        patient_name: str,
        visit_date: str,
        treatment: str,
        amount: str,
        payment_method: str,
    ) -> Path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        patient_folder = self.patient_dir(patient_name)
        bill_path = patient_folder / f"BILL_{timestamp}.pdf"

        c = canvas.Canvas(str(bill_path), pagesize=A4)
        width, height = A4

        c.setFont("Helvetica-Bold", 22)
        c.drawString(20 * mm, height - 25 * mm, CLINIC_NAME)

        c.setFont("Helvetica", 12)
        c.drawString(20 * mm, height - 33 * mm, DOCTOR_NAME)

        c.setLineWidth(0.8)
        c.line(20 * mm, height - 37 * mm, width - 20 * mm, height - 37 * mm)

        y = height - 55 * mm
        c.setFont("Helvetica", 12)
        fields = [
            ("Patient Name", patient_name),
            ("Date", visit_date),
            ("Treatment", treatment),
            ("Amount", f"₹ {amount}"),
            ("Payment Mode", payment_method),
        ]

        for label, value in fields:
            c.drawString(25 * mm, y, f"{label}:")
            c.setFont("Helvetica-Bold", 12)
            c.drawString(65 * mm, y, str(value))
            c.setFont("Helvetica", 12)
            y -= 12 * mm

        c.setLineWidth(0.4)
        c.line(20 * mm, 45 * mm, width - 20 * mm, 45 * mm)
        c.setFont("Helvetica-Oblique", 10)
        c.drawString(20 * mm, 35 * mm, "Thank you for visiting Malathi Dental Clinic.")

        c.showPage()
        c.save()
        return bill_path

    def append_master_record(
        self,
        patient_name: str,
        visit_date: str,
        treatment: str,
        amount: str,
        payment_method: str,
        bill_path: Path,
    ) -> None:
        if not self.records_file.exists():
            self._ensure_structure()

        doc = Document(self.records_file)
        doc.add_paragraph("=" * 60)
        doc.add_paragraph(f"Patient Name: {patient_name}")
        doc.add_paragraph(f"Date: {visit_date}")
        doc.add_paragraph(f"Treatment: {treatment}")
        doc.add_paragraph(f"Amount: ₹ {amount}")
        doc.add_paragraph(f"Payment Method: {payment_method}")
        doc.add_paragraph(f"Bill File: {bill_path}")
        doc.add_paragraph("")
        doc.save(self.records_file)

    def list_all_bills(self) -> list[Path]:
        return sorted(self.patients_dir.glob("**/*.pdf"), key=lambda p: p.stat().st_mtime, reverse=True)

    def search_patient_bills(self, query: str) -> list[Path]:
        query = query.lower().strip()
        if len(query) < 2:
            return []

        matches = []
        for patient_folder in self.patients_dir.iterdir():
            if patient_folder.is_dir() and query in patient_folder.name.lower():
                matches.extend(patient_folder.glob("*.pdf"))
        return sorted(matches, key=lambda p: p.stat().st_mtime, reverse=True)


def open_path(path: Path) -> None:
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.run(["open", str(path)], check=False)
        else:
            subprocess.run(["xdg-open", str(path)], check=False)
    except Exception as exc:
        raise RuntimeError(f"Unable to open: {path}\n{exc}") from exc


class DashboardPage(QWidget):
    def __init__(self, on_new, on_search, on_bills, on_records):
        super().__init__()
        layout = QVBoxLayout(self)
        layout.setSpacing(16)

        title = QLabel(CLINIC_NAME)
        title.setAlignment(Qt.AlignCenter)
        title.setFont(QFont("Segoe UI", 24, QFont.Bold))

        subtitle = QLabel(DOCTOR_NAME)
        subtitle.setAlignment(Qt.AlignCenter)
        subtitle.setFont(QFont("Segoe UI", 12))

        btn_layout = QVBoxLayout()
        btn_layout.setSpacing(14)

        buttons = [
            ("New Patient", on_new),
            ("Search Patient", on_search),
            ("View Bills", on_bills),
            ("Records", on_records),
        ]
        for text, callback in buttons:
            btn = QPushButton(text)
            btn.setMinimumHeight(70)
            btn.setFont(QFont("Segoe UI", 14, QFont.Bold))
            btn.clicked.connect(callback)
            btn_layout.addWidget(btn)

        layout.addStretch(1)
        layout.addWidget(title)
        layout.addWidget(subtitle)
        layout.addSpacing(20)
        layout.addLayout(btn_layout)
        layout.addStretch(1)


class NewPatientPage(QWidget):
    def __init__(self, fm: FileManager, on_back):
        super().__init__()
        self.fm = fm
        self.on_back = on_back

        root = QVBoxLayout(self)

        heading = QLabel("New Patient Entry")
        heading.setFont(QFont("Segoe UI", 18, QFont.Bold))

        form = QFormLayout()
        self.name_input = QLineEdit()
        self.date_input = QDateEdit()
        self.date_input.setDate(QDate.currentDate())
        self.date_input.setDisplayFormat("dd-MM-yyyy")
        self.date_input.setCalendarPopup(True)

        self.treatment_input = QComboBox()
        self.treatment_input.addItems(TREATMENTS)

        self.amount_input = QLineEdit()
        self.amount_input.setPlaceholderText("Enter amount in INR")

        self.payment_input = QComboBox()
        self.payment_input.addItems(PAYMENT_METHODS)

        form.addRow("Patient name", self.name_input)
        form.addRow("Date", self.date_input)
        form.addRow("Treatment", self.treatment_input)
        form.addRow("Amount (₹)", self.amount_input)
        form.addRow("Payment method", self.payment_input)

        button_row = QHBoxLayout()
        save_btn = QPushButton("Save & Generate Bill")
        back_btn = QPushButton("Back")
        save_btn.clicked.connect(self.save_patient)
        back_btn.clicked.connect(self.on_back)
        button_row.addWidget(save_btn)
        button_row.addWidget(back_btn)

        root.addWidget(heading)
        root.addSpacing(10)
        root.addLayout(form)
        root.addSpacing(10)
        root.addLayout(button_row)
        root.addStretch(1)

    def save_patient(self) -> None:
        name = self.name_input.text().strip()
        amount = self.amount_input.text().strip()

        if not name:
            QMessageBox.warning(self, "Validation", "Patient name is required.")
            return

        if not amount.isdigit():
            QMessageBox.warning(self, "Validation", "Amount must be a valid number.")
            return

        visit_date = self.date_input.date().toString("dd-MM-yyyy")
        treatment = self.treatment_input.currentText()
        payment_method = self.payment_input.currentText()

        try:
            bill_path = self.fm.generate_bill_pdf(name, visit_date, treatment, amount, payment_method)
            self.fm.append_master_record(name, visit_date, treatment, amount, payment_method, bill_path)
        except Exception as exc:
            QMessageBox.critical(self, "Error", f"Failed to save patient details.\n{exc}")
            return

        QMessageBox.information(
            self,
            "Saved",
            f"Patient saved successfully.\nBill created at:\n{bill_path}",
        )
        self.name_input.clear()
        self.amount_input.clear()


class BillsListPage(QWidget):
    def __init__(self, fm: FileManager, title: str, on_back, search_mode: bool = False):
        super().__init__()
        self.fm = fm
        self.search_mode = search_mode

        layout = QVBoxLayout(self)
        heading = QLabel(title)
        heading.setFont(QFont("Segoe UI", 18, QFont.Bold))
        layout.addWidget(heading)

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Enter patient name (min 2 characters)")
        self.search_input.textChanged.connect(self.refresh_list)

        if self.search_mode:
            layout.addWidget(self.search_input)

        self.list_widget = QListWidget()
        self.list_widget.itemDoubleClicked.connect(self.open_selected_bill)
        layout.addWidget(self.list_widget)

        back_btn = QPushButton("Back")
        back_btn.clicked.connect(on_back)
        layout.addWidget(back_btn)

    def showEvent(self, event):
        super().showEvent(event)
        self.refresh_list()

    def refresh_list(self) -> None:
        self.list_widget.clear()
        try:
            if self.search_mode:
                query = self.search_input.text()
                bills = self.fm.search_patient_bills(query)
                if len(query.strip()) < 2:
                    self.list_widget.addItem("Type at least 2 characters to search.")
                    return
            else:
                bills = self.fm.list_all_bills()

            if not bills:
                self.list_widget.addItem("No bills found.")
                return

            for bill in bills:
                item = QListWidgetItem(f"{bill.parent.name}  |  {bill.name}")
                item.setData(Qt.UserRole, str(bill))
                self.list_widget.addItem(item)
        except Exception as exc:
            self.list_widget.addItem(f"Unable to load bills: {exc}")

    def open_selected_bill(self, item: QListWidgetItem) -> None:
        path_str = item.data(Qt.UserRole)
        if not path_str:
            return

        bill_path = Path(path_str)
        if not bill_path.exists():
            QMessageBox.warning(self, "Missing file", "Selected bill does not exist.")
            self.refresh_list()
            return

        try:
            open_path(bill_path)
        except RuntimeError as exc:
            QMessageBox.critical(self, "Open failed", str(exc))


class RecordsPage(QWidget):
    def __init__(self, fm: FileManager, on_back):
        super().__init__()
        self.fm = fm

        layout = QVBoxLayout(self)
        heading = QLabel("Records")
        heading.setFont(QFont("Segoe UI", 18, QFont.Bold))

        open_records_btn = QPushButton("Open MASTER_RECORDS.docx")
        open_records_btn.clicked.connect(self.open_records)

        open_patients_btn = QPushButton("Open Patients Folder")
        open_patients_btn.clicked.connect(self.open_patients_folder)

        back_btn = QPushButton("Back")
        back_btn.clicked.connect(on_back)

        layout.addWidget(heading)
        layout.addSpacing(10)
        layout.addWidget(open_records_btn)
        layout.addWidget(open_patients_btn)
        layout.addStretch(1)
        layout.addWidget(back_btn)

    def open_records(self) -> None:
        try:
            self.fm._ensure_structure()
            open_path(self.fm.records_file)
        except Exception as exc:
            QMessageBox.critical(self, "Error", f"Unable to open records file.\n{exc}")

    def open_patients_folder(self) -> None:
        try:
            self.fm._ensure_structure()
            open_path(self.fm.patients_dir)
        except Exception as exc:
            QMessageBox.critical(self, "Error", f"Unable to open patients folder.\n{exc}")


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(CLINIC_NAME)
        self.resize(1200, 800)

        self.fm = FileManager(Path(__file__).resolve().parent)

        self.stack = QStackedWidget()
        self.setCentralWidget(self.stack)

        self.dashboard = DashboardPage(
            on_new=lambda: self.stack.setCurrentWidget(self.new_patient_page),
            on_search=lambda: self.stack.setCurrentWidget(self.search_page),
            on_bills=lambda: self.stack.setCurrentWidget(self.view_bills_page),
            on_records=lambda: self.stack.setCurrentWidget(self.records_page),
        )

        self.new_patient_page = NewPatientPage(self.fm, self.go_home)
        self.search_page = BillsListPage(self.fm, "Search Patient Bills", self.go_home, search_mode=True)
        self.view_bills_page = BillsListPage(self.fm, "All Patient Bills", self.go_home)
        self.records_page = RecordsPage(self.fm, self.go_home)

        self.stack.addWidget(self.dashboard)
        self.stack.addWidget(self.new_patient_page)
        self.stack.addWidget(self.search_page)
        self.stack.addWidget(self.view_bills_page)
        self.stack.addWidget(self.records_page)
        self.stack.setCurrentWidget(self.dashboard)

        self.showMaximized()

    def go_home(self):
        self.stack.setCurrentWidget(self.dashboard)


def apply_dark_theme(app: QApplication) -> None:
    app.setStyleSheet(
        """
        QWidget {
            background-color: #121417;
            color: #E8EDF2;
            font-size: 14px;
        }
        QLabel {
            color: #F0F4F8;
        }
        QPushButton {
            background-color: #1F2A36;
            border: 1px solid #2D3A46;
            border-radius: 10px;
            padding: 10px;
        }
        QPushButton:hover {
            background-color: #2A3947;
        }
        QPushButton:pressed {
            background-color: #18212B;
        }
        QLineEdit, QDateEdit, QComboBox, QListWidget {
            background-color: #1A2028;
            border: 1px solid #2F3C48;
            border-radius: 8px;
            padding: 8px;
            color: #E8EDF2;
        }
        QListWidget::item:selected {
            background-color: #31485E;
        }
        """
    )


def main() -> int:
    app = QApplication(sys.argv)
    apply_dark_theme(app)

    window = MainWindow()
    window.show()

    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
